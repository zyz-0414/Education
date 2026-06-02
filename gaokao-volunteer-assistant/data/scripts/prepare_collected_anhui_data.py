from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Iterator

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_EXTERNAL_ROOT = PROJECT_ROOT.parent / "data"
CLEANED = PROJECT_ROOT / "data" / "cleaned" / "anhui"
SOURCES = PROJECT_ROOT / "data" / "sources" / "source_index.csv"

PROVINCE_CODE = "AH"
BATCH_CODE = "ordinary_undergraduate"
PARSER_VERSION = "collected_anhui_excel_v1"

SCORE_HEADERS = [
    "year",
    "province_code",
    "subject_track",
    "score",
    "count",
    "cumulative_count",
    "rank_min",
    "rank_max",
    "source_id",
]
COLLEGE_HEADERS = [
    "college_code",
    "college_name",
    "province",
    "city",
    "level",
    "ownership",
    "tags",
    "official_site",
    "source_id",
]
MAJOR_HEADERS = [
    "major_code",
    "major_name",
    "major_category",
    "degree_category",
    "duration",
    "notes",
    "source_id",
]
GROUP_HEADERS = [
    "year",
    "province_code",
    "batch_code",
    "subject_track",
    "college_code",
    "group_code",
    "college_name",
    "subject_requirement",
    "group_note",
    "source_id",
]
PLAN_HEADERS = [
    "year",
    "province_code",
    "batch_code",
    "subject_track",
    "college_code",
    "group_code",
    "major_code",
    "major_name",
    "subject_requirement",
    "plan_count",
    "tuition",
    "duration",
    "campus",
    "note",
    "source_id",
]
ADMISSION_HEADERS = [
    "year",
    "province_code",
    "batch_code",
    "subject_track",
    "college_code",
    "group_code",
    "major_code",
    "min_score",
    "min_rank",
    "avg_score",
    "avg_rank",
    "max_score",
    "max_rank",
    "admitted_count",
    "source_id",
]
SOURCE_HEADERS = [
    "source_id",
    "title",
    "source_url",
    "publisher",
    "published_at",
    "fetched_at",
    "file_hash",
    "parser_version",
    "review_status",
    "notes",
]


@dataclass(frozen=True)
class SourceSpec:
    source_id: str
    title: str
    path: Path | None
    publisher: str = "用户收集数据"
    published_at: str = ""
    review_status: str = "collected"
    notes: str = ""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build cleaned Anhui CSV files from the collected Excel data directory.",
    )
    parser.add_argument(
        "--external-root",
        type=Path,
        default=DEFAULT_EXTERNAL_ROOT,
        help="Directory that contains the collected Anhui Excel files.",
    )
    return parser.parse_args()


def norm(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_int(value: object) -> int | None:
    text = norm(value)
    if not text or text in {"-", "--", "—"}:
        return None
    match = re.search(r"\d+", text.replace(",", ""))
    return int(match.group(0)) if match else None


def normalize_code(value: object, width: int | None = None) -> str:
    text = norm(value)
    if not text:
        return ""
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    if width and text.isdigit():
        return text.zfill(width)
    return text


def normalize_subject_track(value: object) -> str:
    text = norm(value)
    if "物理" in text or text == "理科":
        return "physics"
    if "历史" in text or text == "文科":
        return "history"
    raise ValueError(f"unknown subject track: {text}")


def normalize_subject_requirement(value: object) -> str:
    text = norm(value)
    text = re.sub(r"首选\s*(物理|历史)\s*[，,、+ ]*\s*再选\s*", "", text)
    text = text.replace("·", "").replace("(", "").replace(")", "")
    text = text.replace("（", "").replace("）", "").replace("2科必选", "")
    text = text.replace("生物学", "生物")

    if not text or "不限" in text:
        return "不限"

    subjects: list[str] = []
    matchers = [
        ("化学", ("化学", "化")),
        ("生物", ("生物", "生")),
        ("思想政治", ("思想政治", "政治", "政")),
        ("地理", ("地理", "地")),
    ]
    for label, tokens in matchers:
        if any(token in text for token in tokens):
            subjects.append(label)

    return "+".join(subjects) if subjects else "不限"


def normalize_batch(value: object) -> str | None:
    text = norm(value)
    return BATCH_CODE if text in {"本科批", "普通本科批"} else None


def source_id(kind: str, year: int) -> str:
    return f"src_collected_ah_{year}_{kind}"


def sha256_file(path: Path | None) -> str:
    if path is None or not path.exists():
        return ""
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def find_one(root: Path, filename: str) -> Path:
    matches = list(root.rglob(filename))
    if not matches:
        raise FileNotFoundError(f"missing collected data file: {filename} under {root}")
    return matches[0]


def find_header_row(ws: openpyxl.worksheet.worksheet.Worksheet, required: set[str]) -> tuple[int, list[str]]:
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        headers = [norm(value) for value in row]
        if required.issubset(set(headers)):
            return row_index, headers
    raise ValueError(f"header row not found in {ws.title}; required={sorted(required)}")


def iter_dict_rows(path: Path, required_headers: set[str]) -> Iterator[dict[str, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            if worksheet.max_row < 2:
                continue
            try:
                header_row, headers = find_header_row(worksheet, required_headers)
            except ValueError:
                continue

            empty_streak = 0
            for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                values = [norm(value) for value in row]
                if not any(values):
                    empty_streak += 1
                    if empty_streak >= 200:
                        break
                    continue
                empty_streak = 0
                yield {
                    header: values[index] if index < len(values) else ""
                    for index, header in enumerate(headers)
                    if header
                }
    finally:
        workbook.close()


def collect_score_segments(path: Path, year: int, subject_track: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = None
        header_row = 0
        for candidate in workbook.worksheets:
            try:
                header_row, _headers = find_header_row(candidate, {"分数", "人数", "累计人数"})
                worksheet = candidate
                break
            except ValueError:
                continue
        if worksheet is None:
            raise ValueError(f"score segment sheet not found in {path}")
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            score = parse_int(row[0] if len(row) > 0 else None)
            count = parse_int(row[1] if len(row) > 1 else None)
            cumulative = parse_int(row[2] if len(row) > 2 else None)
            if score is None or count is None or cumulative is None:
                continue
            rows.append(
                {
                    "year": year,
                    "province_code": PROVINCE_CODE,
                    "subject_track": subject_track,
                    "score": score,
                    "count": count,
                    "cumulative_count": cumulative,
                    "rank_min": cumulative - count + 1,
                    "rank_max": cumulative,
                    "source_id": source_id("score_segments", year),
                }
            )
    finally:
        workbook.close()
    rows.sort(key=lambda row: (row["year"], row["subject_track"], -int(row["score"])))
    return rows


def group_key(row: dict[str, object]) -> tuple[object, ...]:
    return (
        row["year"],
        row["province_code"],
        row["batch_code"],
        row["subject_track"],
        row["college_code"],
        row["group_code"],
    )


def add_group(
    groups: dict[tuple[object, ...], dict[str, object]],
    row: dict[str, object],
    college_name: str,
    subject_requirement: str,
    note: str,
    row_source_id: str,
) -> None:
    key = group_key(row)
    existing = groups.get(key)
    if existing:
        if not existing.get("subject_requirement") or existing["subject_requirement"] == "不限":
            existing["subject_requirement"] = subject_requirement
        if note not in str(existing.get("group_note", "")):
            existing["group_note"] = f"{existing.get('group_note', '')}; {note}".strip("; ")
        return

    groups[key] = {
        "year": row["year"],
        "province_code": row["province_code"],
        "batch_code": row["batch_code"],
        "subject_track": row["subject_track"],
        "college_code": row["college_code"],
        "group_code": row["group_code"],
        "college_name": college_name,
        "subject_requirement": subject_requirement,
        "group_note": note,
        "source_id": row_source_id,
    }


def make_college(
    college_code: str,
    college_name: str,
    source: str,
    province: str = "",
    city: str = "",
    ownership: str = "",
) -> dict[str, object]:
    tags = ["普通本科批", "专业组数据"]
    return {
        "college_code": college_code,
        "college_name": college_name,
        "province": province,
        "city": city,
        "level": "本科",
        "ownership": ownership,
        "tags": json.dumps(tags, ensure_ascii=False),
        "official_site": "",
        "source_id": source,
    }


def collect_2025_plans(
    path: Path,
    groups: dict[tuple[object, ...], dict[str, object]],
    colleges: dict[str, dict[str, object]],
    majors: dict[tuple[str, str, str], dict[str, object]],
) -> list[dict[str, object]]:
    rows: dict[tuple[object, ...], dict[str, object]] = {}
    sid = source_id("enrollment_plan", 2025)
    for item in iter_dict_rows(path, {"年份", "生源地", "科类", "批次", "院校代码", "专业组代码", "专业代码"}):
        if normalize_batch(item.get("批次")) != BATCH_CODE:
            continue
        year = parse_int(item.get("年份"))
        if year != 2025:
            continue

        subject_track = normalize_subject_track(item.get("科类"))
        college_code = normalize_code(item.get("院校代码"), 4)
        group_code = normalize_code(item.get("专业组代码"), 3)
        major_code = normalize_code(item.get("专业代码"), 2)
        major_name = norm(item.get("专业名称"))
        college_name = norm(item.get("院校名称"))
        plan_count = parse_int(item.get("计划人数"))
        if not all([college_code, group_code, major_code, major_name, college_name]) or plan_count is None:
            continue

        requirement = normalize_subject_requirement(item.get("选科要求"))
        row = {
            "year": 2025,
            "province_code": PROVINCE_CODE,
            "batch_code": BATCH_CODE,
            "subject_track": subject_track,
            "college_code": college_code,
            "group_code": group_code,
            "major_code": major_code,
            "major_name": major_name,
            "subject_requirement": requirement,
            "plan_count": plan_count,
            "tuition": parse_int(item.get("学费")) or "",
            "duration": "四年",
            "campus": "",
            "note": norm(item.get("专业备注")),
            "source_id": sid,
        }
        key = group_key(row) + (major_code,)
        if key in rows:
            rows[key]["plan_count"] = int(rows[key]["plan_count"]) + plan_count
        else:
            rows[key] = row

        colleges[college_code] = make_college(college_code, college_name, sid)
        majors.setdefault(
            (major_code, major_name, sid),
            {
                "major_code": major_code,
                "major_name": major_name,
                "major_category": "",
                "degree_category": "本科",
                "duration": "四年",
                "notes": "安徽2025招生计划省编专业代码，非国家标准专业代码",
                "source_id": sid,
            },
        )
        add_group(groups, row, college_name, requirement, "2025招生计划", sid)
    return list(rows.values())


def collect_2024_plans(
    path: Path,
    groups: dict[tuple[object, ...], dict[str, object]],
    colleges: dict[str, dict[str, object]],
    majors: dict[tuple[str, str, str], dict[str, object]],
) -> list[dict[str, object]]:
    rows: dict[tuple[object, ...], dict[str, object]] = {}
    sid = source_id("enrollment_plan", 2024)
    for item in iter_dict_rows(path, {"年份", "省份", "科类", "批次", "招生代码", "专业组编号", "专业代码"}):
        if normalize_batch(item.get("批次")) != BATCH_CODE:
            continue
        year = parse_int(item.get("年份"))
        if year != 2024:
            continue

        subject_track = normalize_subject_track(item.get("科类"))
        college_code = normalize_code(item.get("招生代码"), 4)
        group_code = normalize_code(item.get("专业组编号"), 3)
        major_code = normalize_code(item.get("专业代码"), 2)
        major_name = norm(item.get("专业名称"))
        college_name = norm(item.get("院校名称"))
        plan_count = parse_int(item.get("招生计划人数"))
        if not all([college_code, group_code, major_code, major_name, college_name]) or plan_count is None:
            continue

        requirement = normalize_subject_requirement(item.get("报考要求"))
        row = {
            "year": 2024,
            "province_code": PROVINCE_CODE,
            "batch_code": BATCH_CODE,
            "subject_track": subject_track,
            "college_code": college_code,
            "group_code": group_code,
            "major_code": major_code,
            "major_name": major_name,
            "subject_requirement": requirement,
            "plan_count": plan_count,
            "tuition": parse_int(item.get("学费")) or "",
            "duration": "四年",
            "campus": "",
            "note": norm(item.get("专业备注")),
            "source_id": sid,
        }
        key = group_key(row) + (major_code,)
        if key in rows:
            rows[key]["plan_count"] = int(rows[key]["plan_count"]) + plan_count
        else:
            rows[key] = row

        colleges.setdefault(college_code, make_college(college_code, college_name, sid))
        majors.setdefault(
            (major_code, major_name, sid),
            {
                "major_code": major_code,
                "major_name": major_name,
                "major_category": "",
                "degree_category": "本科",
                "duration": "四年",
                "notes": "安徽2024招生计划省编专业代码，非国家标准专业代码",
                "source_id": sid,
            },
        )
        add_group(groups, row, college_name, requirement, "2024招生计划", sid)
    return list(rows.values())


def collect_2025_admissions(
    path: Path,
    groups: dict[tuple[object, ...], dict[str, object]],
    colleges: dict[str, dict[str, object]],
) -> list[dict[str, object]]:
    rows: dict[tuple[object, ...], dict[str, object]] = {}
    sid = source_id("group_admission", 2025)
    for item in iter_dict_rows(path, {"年份", "生源地", "批次", "院校代码", "科类", "专业组代码", "最低分"}):
        if normalize_batch(item.get("批次")) != BATCH_CODE:
            continue
        year = parse_int(item.get("年份"))
        if year != 2025:
            continue

        subject_track = normalize_subject_track(item.get("科类"))
        row = {
            "year": 2025,
            "province_code": PROVINCE_CODE,
            "batch_code": BATCH_CODE,
            "subject_track": subject_track,
            "college_code": normalize_code(item.get("院校代码"), 4),
            "group_code": normalize_code(item.get("专业组代码"), 3),
            "major_code": "",
            "min_score": parse_int(item.get("最低分")) or "",
            "min_rank": parse_int(item.get("最低位次")) or "",
            "avg_score": "",
            "avg_rank": "",
            "max_score": "",
            "max_rank": "",
            "admitted_count": parse_int(item.get("投档人数")) or "",
            "source_id": sid,
        }
        if not row["college_code"] or not row["group_code"] or not row["min_rank"]:
            continue
        key = group_key(row)
        rows[key] = row
        college_name = norm(item.get("院校名称"))
        requirement = normalize_subject_requirement(item.get("选科要求"))
        colleges.setdefault(str(row["college_code"]), make_college(str(row["college_code"]), college_name, sid))
        add_group(groups, row, college_name, requirement, "2025专业组投档线", sid)
    return list(rows.values())


def parse_2024_group(value: object) -> tuple[str, str] | None:
    text = norm(value)
    match = re.search(r"(\d{4})\s*[（(]\s*[A-Z]?(\d{3})\s*[）)]", text)
    if not match:
        return None
    return match.group(1), match.group(2)


def collect_2024_admissions(
    path: Path,
    groups: dict[tuple[object, ...], dict[str, object]],
    colleges: dict[str, dict[str, object]],
) -> list[dict[str, object]]:
    rows: dict[tuple[object, ...], dict[str, object]] = {}
    sid = source_id("group_admission", 2024)
    for item in iter_dict_rows(path, {"年份", "学校", "科类", "批次", "专业组", "最低分", "最低分段"}):
        if normalize_batch(item.get("批次")) != BATCH_CODE or norm(item.get("类别")) != "普通类":
            continue
        year = parse_int(item.get("年份"))
        parsed_group = parse_2024_group(item.get("专业组"))
        if year != 2024 or parsed_group is None:
            continue

        college_code, group_code = parsed_group
        subject_track = normalize_subject_track(item.get("科类"))
        row = {
            "year": 2024,
            "province_code": PROVINCE_CODE,
            "batch_code": BATCH_CODE,
            "subject_track": subject_track,
            "college_code": college_code,
            "group_code": group_code,
            "major_code": "",
            "min_score": parse_int(item.get("最低分")) or "",
            "min_rank": parse_int(item.get("最低分段")) or "",
            "avg_score": "",
            "avg_rank": "",
            "max_score": "",
            "max_rank": "",
            "admitted_count": "",
            "source_id": sid,
        }
        if not row["min_rank"]:
            continue
        key = group_key(row)
        rows[key] = row
        college_name = norm(item.get("学校"))
        requirement = normalize_subject_requirement(item.get("选科要求"))
        colleges.setdefault(
            college_code,
            make_college(
                college_code,
                college_name,
                sid,
                province=norm(item.get("省")),
                city=norm(item.get("城市")),
                ownership=norm(item.get("办学性质")),
            ),
        )
        add_group(groups, row, college_name, requirement, "2024专业组投档线", sid)
    return list(rows.values())


def enrich_colleges_from_2025_major_admissions(path: Path, colleges: dict[str, dict[str, object]]) -> None:
    for item in iter_dict_rows(path, {"年份", "生源地", "批次", "科类", "院校代码", "院校名称", "最低位次"}):
        if parse_int(item.get("年份")) != 2025 or normalize_batch(item.get("批次")) != BATCH_CODE:
            continue
        college_code = normalize_code(item.get("院校代码"), 4)
        if not college_code or college_code not in colleges:
            continue
        row = colleges[college_code]
        row["province"] = row.get("province") or norm(item.get("所在省"))
        row["city"] = row.get("city") or norm(item.get("城市"))
        row["ownership"] = row.get("ownership") or norm(item.get("公私性质"))


def build_charter_rules() -> list[dict[str, object]]:
    return []


def build_sources(specs: Iterable[SourceSpec]) -> list[dict[str, object]]:
    fetched_at = date.today().isoformat()
    rows = [
        {
            "source_id": "src_ah_exam",
            "title": "安徽省教育招生考试院",
            "source_url": "https://www.ahzsks.cn/",
            "publisher": "安徽省教育招生考试院",
            "published_at": "",
            "fetched_at": fetched_at,
            "file_hash": "",
            "parser_version": "",
            "review_status": "planned",
            "notes": "考试院入口，政策表外键保留",
        },
        {
            "source_id": "src_ah_service",
            "title": "安徽省普通高校招生考生服务平台",
            "source_url": "https://xgk.ahzsks.cn/",
            "publisher": "安徽省教育招生考试院",
            "published_at": "",
            "fetched_at": fetched_at,
            "file_hash": "",
            "parser_version": "",
            "review_status": "planned",
            "notes": "志愿填报服务平台入口",
        },
        {
            "source_id": "src_chsi",
            "title": "阳光高考平台",
            "source_url": "https://gaokao.chsi.com.cn/",
            "publisher": "教育部阳光高考",
            "published_at": "",
            "fetched_at": fetched_at,
            "file_hash": "",
            "parser_version": "",
            "review_status": "planned",
            "notes": "招生计划和章程参考入口",
        },
    ]

    for spec in specs:
        rows.append(
            {
                "source_id": spec.source_id,
                "title": spec.title,
                "source_url": f"local:{spec.path}" if spec.path else "local:collected-data",
                "publisher": spec.publisher,
                "published_at": spec.published_at,
                "fetched_at": fetched_at,
                "file_hash": sha256_file(spec.path),
                "parser_version": PARSER_VERSION,
                "review_status": spec.review_status,
                "notes": spec.notes,
            }
        )
    return rows


def sort_rows(rows: Iterable[dict[str, object]], fields: list[str]) -> list[dict[str, object]]:
    return sorted(rows, key=lambda row: tuple(str(row.get(field, "")) for field in fields))


def main() -> None:
    args = parse_args()
    external_root = args.external_root.resolve()
    if not external_root.exists():
        raise SystemExit(f"external data directory not found: {external_root}")

    files = {
        "score_2024_physics": find_one(external_root, "安徽_一分一段_2024_物理组.xlsx"),
        "score_2024_history": find_one(external_root, "安徽_一分一段_2024_历史组.xlsx"),
        "score_2025_physics": find_one(external_root, "安徽2025一分一段表（物理）.xlsx"),
        "score_2025_history": find_one(external_root, "安徽2025一分一段表（历史）.xlsx"),
        "plan_2024": find_one(external_root, "安徽_招生计划_2024.xlsx"),
        "plan_2025": find_one(external_root, "安徽-2025-招生计划.xlsx"),
        "admission_2024": find_one(external_root, "安徽_投档线_2024.xlsx"),
        "admission_2025": find_one(external_root, "安徽25年专业组投档线最新.xlsx"),
        "major_admission_2025": find_one(external_root, "安徽省2025年专业分数线.xlsx"),
    }

    groups: dict[tuple[object, ...], dict[str, object]] = {}
    colleges: dict[str, dict[str, object]] = {}
    majors: dict[tuple[str, str, str], dict[str, object]] = {}

    score_rows = []
    score_rows.extend(collect_score_segments(files["score_2024_physics"], 2024, "physics"))
    score_rows.extend(collect_score_segments(files["score_2024_history"], 2024, "history"))
    score_rows.extend(collect_score_segments(files["score_2025_physics"], 2025, "physics"))
    score_rows.extend(collect_score_segments(files["score_2025_history"], 2025, "history"))

    plan_rows = []
    plan_rows.extend(collect_2024_plans(files["plan_2024"], groups, colleges, majors))
    plan_rows.extend(collect_2025_plans(files["plan_2025"], groups, colleges, majors))

    admission_rows = []
    admission_rows.extend(collect_2024_admissions(files["admission_2024"], groups, colleges))
    admission_rows.extend(collect_2025_admissions(files["admission_2025"], groups, colleges))
    enrich_colleges_from_2025_major_admissions(files["major_admission_2025"], colleges)

    specs = [
        SourceSpec(source_id("score_segments", 2024), "安徽2024一分一段表（物理/历史）", files["score_2024_physics"], notes="同目录历史组文件一并解析"),
        SourceSpec(source_id("score_segments", 2025), "安徽2025一分一段表（物理/历史）", files["score_2025_physics"], notes="同目录历史组文件一并解析"),
        SourceSpec(source_id("enrollment_plan", 2024), "安徽2024普通本科批招生计划", files["plan_2024"]),
        SourceSpec(source_id("enrollment_plan", 2025), "安徽2025普通本科批招生计划", files["plan_2025"]),
        SourceSpec(source_id("group_admission", 2024), "安徽2024普通本科批专业组投档线", files["admission_2024"]),
        SourceSpec(source_id("group_admission", 2025), "安徽2025普通本科批专业组投档线", files["admission_2025"]),
    ]

    write_csv(CLEANED / "score_segments.csv", SCORE_HEADERS, sort_rows(score_rows, ["year", "subject_track", "score"]))
    write_csv(CLEANED / "colleges.csv", COLLEGE_HEADERS, sort_rows(colleges.values(), ["college_code"]))
    write_csv(CLEANED / "majors.csv", MAJOR_HEADERS, sort_rows(majors.values(), ["major_code", "major_name"]))
    write_csv(CLEANED / "college_groups.csv", GROUP_HEADERS, sort_rows(groups.values(), ["year", "subject_track", "college_code", "group_code"]))
    write_csv(CLEANED / "enrollment_plans.csv", PLAN_HEADERS, sort_rows(plan_rows, ["year", "subject_track", "college_code", "group_code", "major_code"]))
    write_csv(CLEANED / "admission_results.csv", ADMISSION_HEADERS, sort_rows(admission_rows, ["year", "subject_track", "college_code", "group_code"]))
    write_csv(CLEANED / "charter_rules.csv", ["year", "college_code", "scope_type", "scope_code", "rule_type", "condition_json", "original_text", "review_status", "source_id"], build_charter_rules())
    write_csv(SOURCES, SOURCE_HEADERS, build_sources(specs))

    print(
        "built collected Anhui data "
        f"score_segments={len(score_rows)} "
        f"college_groups={len(groups)} "
        f"enrollment_plans={len(plan_rows)} "
        f"admission_results={len(admission_rows)} "
        f"colleges={len(colleges)} "
        f"majors={len(majors)}"
    )


if __name__ == "__main__":
    main()
